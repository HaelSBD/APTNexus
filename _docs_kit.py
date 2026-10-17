"""Shared formatting for APTNexus official docs — mirrors 3RKU style (read-only reference)."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FIGS = Path(r"C:\Users\Jcele\Desktop\Official Documentations\APTNexus\figures")
OUT = Path(r"C:\Users\Jcele\Desktop\Official Documentations\APTNexus")

NAVY = RGBColor(0x0F, 0x23, 0x40)


def set_run_font(run, size=12, bold=False, italic=False, name="Times New Roman", color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_after = Pt(8)

    for style_name, size in [("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12.5), ("Title", 30)]:
        try:
            st = doc.styles[style_name]
        except KeyError:
            continue
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = NAVY
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def new_doc():
    doc = Document()
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.2)
        section.top_margin = Cm(2.3)
        section.bottom_margin = Cm(2.0)
    configure_styles(doc)
    return doc


def add_title_block(doc, title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    set_run_font(r, size=28, bold=True, color=NAVY)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitle)
    set_run_font(r2, size=14, italic=False, color=NAVY)
    doc.add_paragraph()


def add_h1(doc, text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        set_run_font(r, size=20, bold=True, color=NAVY)
    return p


def add_h2(doc, text):
    p = doc.add_heading(text, level=2)
    for r in p.runs:
        set_run_font(r, size=15, bold=True, color=NAVY)
    return p


def add_h3(doc, text):
    p = doc.add_heading(text, level=3)
    for r in p.runs:
        set_run_font(r, size=12.5, bold=True, color=NAVY)
    return p


def add_p(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, size=12)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    r = p.add_run(text)
    set_run_font(r, size=12)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.clear()
    r = p.add_run(text)
    set_run_font(r, size=12)
    return p


def shade_cell(cell, hex_color="D9E8F5"):
    tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, size=11):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold)


def add_kv_table(doc, rows, header=None):
    data = rows
    if header:
        data = [header] + rows
    table = doc.add_table(rows=len(data), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (a, b) in enumerate(data):
        set_cell_text(table.rows[i].cells[0], str(a), bold=(i == 0 and header is not None) or False)
        set_cell_text(table.rows[i].cells[1], str(b), bold=(i == 0 and header is not None))
        if i == 0 and header is not None:
            shade_cell(table.rows[i].cells[0])
            shade_cell(table.rows[i].cells[1])
    doc.add_paragraph()
    return table


def add_grid_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=11)
        shade_cell(table.rows[0].cells[i])
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            set_cell_text(table.rows[ri + 1].cells[ci], str(val), size=10.5)
    doc.add_paragraph()
    return table


def add_figure(doc, filename, caption, width=6.2):
    path = FIGS / filename
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_run_font(r, size=10.5, italic=True)
    doc.add_paragraph()


def add_toc(doc, entries):
    """entries: list of (level, title, page_str)"""
    add_h1(doc, "Contents")
    for level, title, page in entries:
        p = doc.add_paragraph()
        tab_stops = p.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Inches(6.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        indent = 0 if level == 1 else (0.25 if level == 2 else 0.5)
        p.paragraph_format.left_indent = Inches(indent)
        r = p.add_run(f"{title}\t{page}")
        set_run_font(r, size=12, bold=(level == 1))
    doc.add_paragraph()


def add_doc_control(doc, product, prepared_for, title, related, version="1.0", date="14 October 2026"):
    add_h1(doc, "Document control")
    add_kv_table(
        doc,
        [
            ("Product", product),
            ("Prepared for", prepared_for),
        ],
    )
    add_grid_table(
        doc,
        ["Field", "Detail"],
        [
            ["Document title", title],
            ["Version", version],
            ["Date", date],
            ["Related documents", related],
            ["Produced by", "Satok Bridge Digital Sdn. Bhd."],
            ["Classification", "Internal — Official documentation"],
            ["Status", "Approved for programme use"],
        ],
    )


def page_break(doc):
    doc.add_page_break()
